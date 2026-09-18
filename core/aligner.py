"""
Módulo para alinhamento e padronização tipográfica de textos e títulos no Word (.docx).
Garante justificação uniforme, entrelinhas consistente, controle de orfandade de títulos
e formatação por páginas.
"""

from typing import Dict, Any
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from core.config import FormattingConfig
from core.image_handler import _is_image_paragraph


def _get_alignment_enum(align_str: str) -> WD_ALIGN_PARAGRAPH:
    """Converte string de alinhamento para o Enum do python-docx."""
    mapping = {
        "JUSTIFY": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "LEFT": WD_ALIGN_PARAGRAPH.LEFT,
        "CENTER": WD_ALIGN_PARAGRAPH.CENTER,
        "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT,
    }
    return mapping.get(align_str.upper(), WD_ALIGN_PARAGRAPH.JUSTIFY)


def _is_heading(p) -> tuple[bool, int]:
    """
    Identifica se um parágrafo é um título (Heading 1, 2, 3...) pelo estilo ou propriedades.
    Retorna (is_heading, heading_level).
    """
    style_name = (p.style.name or "").lower()
    
    # Título 1 / Heading 1
    if any(k in style_name for k in ["heading 1", "título 1", "titulo 1"]):
        return True, 1
    if any(k in style_name for k in ["heading 2", "título 2", "titulo 2"]):
        return True, 2
    if any(k in style_name for k in ["heading 3", "título 3", "titulo 3"]):
        return True, 3
    if any(k in style_name for k in ["heading", "título", "titulo", "title"]):
        return True, 1

    # Verificar nível no XML (outline level)
    outline = p._p.xpath('.//w:outlineLvl')
    if outline:
        try:
            val = int(outline[0].get(qn('w:val'), '0'))
            return True, val + 1
        except (ValueError, TypeError):
            pass

    return False, 0


def _is_list_item(p) -> bool:
    """Verifica se o parágrafo faz parte de uma lista com marcadores ou numeração."""
    style_name = (p.style.name or "").lower()
    if "list" in style_name or "marcador" in style_name or "bullet" in style_name:
        return True
    num_pr = p._p.xpath('.//w:numPr')
    return len(num_pr) > 0


def _apply_font_to_run(run, font_name: str, font_size_pt: float = None) -> None:
    """Aplica o nome da fonte garantindo compatibilidade no XML de idiomas do Word."""
    if font_name:
        run.font.name = font_name
        rPr = run._r.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)
        rFonts.set(qn('w:cs'), font_name)

    if font_size_pt is not None and font_size_pt > 0:
        run.font.size = Pt(font_size_pt)


def format_text_and_headings(doc: Document, config: FormattingConfig) -> Dict[str, Any]:
    """
    Aplica alinhamento de texto, entrelinhas, recuos e controle de paginação aos parágrafos.
    """
    body_aligned_count = 0
    headings_adjusted_count = 0

    body_alignment = _get_alignment_enum(config.body_alignment)
    heading_alignment = _get_alignment_enum(config.heading_alignment)

    for p in doc.paragraphs:
        # Pular se for parágrafo que contém imagem (já tratado pelo image_handler)
        if _is_image_paragraph(p):
            continue

        text = p.text.strip()
        if not text:
            continue

        is_hd, level = _is_heading(p)
        is_list = _is_list_item(p)

        # 1. Tratar Títulos e Cabeçalhos
        if is_hd:
            p.alignment = heading_alignment
            p.paragraph_format.widow_control = config.widow_control
            
            # Essencial para ajuste de páginas: impede título órfão isolado no fim da página!
            if config.keep_headings_with_next:
                p.paragraph_format.keep_with_next = True

            # Títulos não devem ter recuo de primeira linha
            p.paragraph_format.first_line_indent = Cm(0)
            
            # Espaçamento de títulos
            if level == 1:
                p.paragraph_format.space_before = Pt(12.0)
                p.paragraph_format.space_after = Pt(6.0)
                target_size = config.heading1_size_pt
            elif level == 2:
                p.paragraph_format.space_before = Pt(10.0)
                p.paragraph_format.space_after = Pt(4.0)
                target_size = config.heading2_size_pt
            else:
                p.paragraph_format.space_before = Pt(8.0)
                p.paragraph_format.space_after = Pt(3.0)
                target_size = config.heading3_size_pt

            # Padronização de fonte do título
            if config.change_font:
                for run in p.runs:
                    _apply_font_to_run(run, config.heading_font_name, target_size)

            headings_adjusted_count += 1

        # 2. Tratar Itens de Lista (Marcadores / Números)
        elif is_list:
            # Listas geralmente ficam alinhadas à esquerda ou justificadas sem recuo de primeira linha extra
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.widow_control = config.widow_control
            p.paragraph_format.line_spacing = config.line_spacing
            p.paragraph_format.space_after = Pt(max(2.0, config.space_after_pt * 0.5))

            if config.change_font:
                for run in p.runs:
                    _apply_font_to_run(run, config.font_name, config.font_size_pt)

            body_aligned_count += 1

        # 3. Tratar Corpo de Texto Comum
        else:
            p.alignment = body_alignment
            p.paragraph_format.widow_control = config.widow_control
            p.paragraph_format.line_spacing = config.line_spacing
            p.paragraph_format.space_before = Pt(config.space_before_pt)
            p.paragraph_format.space_after = Pt(config.space_after_pt)

            # Recuo da primeira linha (ex: 1.25 cm se ABNT, 0 se Corporativo)
            if config.first_line_indent_cm > 0:
                p.paragraph_format.first_line_indent = Cm(config.first_line_indent_cm)
            else:
                p.paragraph_format.first_line_indent = Cm(0)

            # Padronização de fonte do corpo de texto
            if config.change_font:
                for run in p.runs:
                    _apply_font_to_run(run, config.font_name, config.font_size_pt)

            body_aligned_count += 1

    # Formatar textos dentro de células de tabelas
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if _is_image_paragraph(p):
                        continue
                    p.paragraph_format.widow_control = config.widow_control
                    p.paragraph_format.line_spacing = 1.0  # tabelas ficam melhores com espaçamento 1.0
                    p.paragraph_format.space_after = Pt(2.0)
                    p.paragraph_format.first_line_indent = Cm(0)
                    if config.change_font:
                        for run in p.runs:
                            _apply_font_to_run(run, config.font_name, max(9.0, config.font_size_pt - 1.0))

    return {
        "body_paragraphs_aligned": body_aligned_count,
        "headings_adjusted": headings_adjusted_count,
        "font_applied": config.font_name if config.change_font else "Original",
        "alignment_applied": config.body_alignment,
    }
