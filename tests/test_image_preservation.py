"""
Teste rigoroso de preservação absoluta de imagens e dados durante a redução de tamanho.
Garante que nenhuma imagem e nenhum dado seja excluído ou corrompido em PDFs e DOCXs.
"""

import os
import io
import unittest
from PIL import Image, ImageDraw
import pymupdf
import docx
import zipfile

from core.compressor import FileCompressor, CompressionConfig


class TestImagePreservation(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.test_dir = os.path.join(os.path.dirname(__file__), "temp_preservation_tests")
        os.makedirs(cls.test_dir, exist_ok=True)

        # 1. Gerar imagens com diferentes características
        # Imagem 1: JPEG foto grande (2200x1500)
        cls.img_jpg_path = os.path.join(cls.test_dir, "foto_grande.jpg")
        img1 = Image.new("RGB", (2200, 1500), color=(60, 130, 210))
        d1 = ImageDraw.Draw(img1)
        for x in range(0, 2200, 80):
            d1.line([(x, 0), (x, 1500)], fill=(255, 255, 255), width=3)
        img1.save(cls.img_jpg_path, format="JPEG", quality=95)

        # Imagem 2: PNG com transparência alfa (800x600)
        cls.img_png_alpha_path = os.path.join(cls.test_dir, "grafico_alpha.png")
        img2 = Image.new("RGBA", (800, 600), (240, 100, 50, 160))
        d2 = ImageDraw.Draw(img2)
        d2.ellipse([50, 50, 750, 550], fill=(50, 200, 100, 200), outline=(255, 255, 255), width=6)
        img2.save(cls.img_png_alpha_path, format="PNG")

        # Imagem 3: PNG em tons de cinza (500x500)
        cls.img_png_gray_path = os.path.join(cls.test_dir, "grafico_cinza.png")
        img3 = Image.new("L", (500, 500), color=180)
        img3.save(cls.img_png_gray_path, format="PNG")

        # 2. Criar PDF multipágina com todas as imagens e textos
        cls.pdf_path = os.path.join(cls.test_dir, "documento_com_todas_imagens.pdf")
        doc_pdf = pymupdf.open()
        
        # Página 1: Texto + Imagem JPEG
        p1 = doc_pdf.new_page(width=595, height=842)
        p1.insert_text((50, 50), "PÁGINA 1: Relatório com Imagem JPEG", fontsize=14)
        with open(cls.img_jpg_path, "rb") as f:
            p1.insert_image(pymupdf.Rect(50, 80, 545, 400), stream=f.read())
        p1.insert_text((50, 420), "Texto descritivo abaixo da imagem 1.", fontsize=11)

        # Página 2: Texto + Imagem PNG Alfa + Imagem Cinza
        p2 = doc_pdf.new_page(width=595, height=842)
        p2.insert_text((50, 50), "PÁGINA 2: Gráficos PNG com e sem transparência", fontsize=14)
        with open(cls.img_png_alpha_path, "rb") as f:
            p2.insert_image(pymupdf.Rect(50, 80, 350, 300), stream=f.read())
        with open(cls.img_png_gray_path, "rb") as f:
            p2.insert_image(pymupdf.Rect(50, 320, 250, 500), stream=f.read())
        p2.insert_text((50, 520), "Texto de rodapé da página 2 que nunca pode ser apagado.", fontsize=11)

        doc_pdf.save(cls.pdf_path)
        doc_pdf.close()

        # 3. Criar DOCX com todas as imagens e parágrafos
        cls.docx_path = os.path.join(cls.test_dir, "documento_com_todas_imagens.docx")
        doc_word = docx.Document()
        doc_word.add_heading("Documento com Múltiplas Imagens para Validação", 0)
        doc_word.add_paragraph("Texto inicial do parágrafo 1.")
        doc_word.add_picture(cls.img_jpg_path)
        doc_word.add_paragraph("Texto intermediário entre imagens.")
        doc_word.add_picture(cls.img_png_alpha_path)
        doc_word.add_paragraph("Texto final que não pode sumir.")
        doc_word.add_picture(cls.img_png_gray_path)
        doc_word.save(cls.docx_path)

    @classmethod
    def tearDownClass(cls):
        if os.path.exists(cls.test_dir):
            for f in os.listdir(cls.test_dir):
                try:
                    os.remove(os.path.join(cls.test_dir, f))
                except Exception:
                    pass
            try:
                os.rmdir(cls.test_dir)
            except Exception:
                pass

    def test_pdf_preserves_every_single_image(self):
        """Verifica se 100% das imagens do PDF permanecem intactas e visíveis."""
        compressor = FileCompressor()
        out_pdf = os.path.join(self.test_dir, "documento_com_todas_imagens_reduzido.pdf")

        # Contagem de imagens e páginas no PDF de entrada
        doc_before = pymupdf.open(self.pdf_path)
        pages_before = len(doc_before)
        images_before = sum(len(p.get_images()) for p in doc_before)
        doc_before.close()

        result = compressor.compress_file(self.pdf_path, output_path=out_pdf)
        self.assertTrue(result.success, f"Falha na compressão do PDF: {result.error_message}")
        self.assertTrue(os.path.exists(out_pdf))

        # Inspecionar o PDF de saída
        doc_after = pymupdf.open(out_pdf)
        pages_after = len(doc_after)
        images_after = sum(len(p.get_images()) for p in doc_after)

        # 1. Nenhuma página pode sumir
        self.assertEqual(pages_after, pages_before, "O número de páginas do PDF foi alterado!")

        # 2. NENHUMA IMAGEM PODE SUMIR
        self.assertEqual(images_after, images_before, f"Imagens foram perdidas! Antes: {images_before}, Depois: {images_after}")

        # 3. Teste de renderização: todas as páginas devem renderizar perfeitamente
        for idx, page in enumerate(doc_after):
            pix = page.get_pixmap(dpi=72)
            self.assertGreater(pix.width, 0)
            self.assertGreater(pix.height, 0)
            # O texto da página deve continuar legível
            text = page.get_text()
            if idx == 0:
                self.assertIn("PÁGINA 1", text)
                self.assertIn("Texto descritivo abaixo da imagem 1.", text)
            elif idx == 1:
                self.assertIn("PÁGINA 2", text)
                self.assertIn("Texto de rodapé", text)

        doc_after.close()

    def test_docx_preserves_every_single_image(self):
        """Verifica se 100% das imagens do DOCX permanecem intactas no arquivo."""
        compressor = FileCompressor()
        out_docx = os.path.join(self.test_dir, "documento_com_todas_imagens_reduzido.docx")

        # Inspecionar imagens dentro do zip original
        with zipfile.ZipFile(self.docx_path, "r") as z_orig:
            orig_media = [f for f in z_orig.namelist() if "media/" in f.lower()]

        result = compressor.compress_file(self.docx_path, output_path=out_docx)
        self.assertTrue(result.success, f"Falha na compressão do DOCX: {result.error_message}")
        self.assertTrue(os.path.exists(out_docx))

        # Inspecionar imagens dentro do zip comprimido
        with zipfile.ZipFile(out_docx, "r") as z_out:
            comp_media = [f for f in z_out.namelist() if "media/" in f.lower()]
            # 1. Cada arquivo de imagem que existia continua existindo
            self.assertEqual(len(comp_media), len(orig_media), f"Imagens no DOCX foram perdidas! Antes: {len(orig_media)}, Depois: {len(comp_media)}")
            self.assertEqual(set(comp_media), set(orig_media), "Nomes de arquivos de imagem no pacote DOCX divergiram!")

            # 2. Cada imagem abre perfeitamente com Pillow
            for f in comp_media:
                data = z_out.read(f)
                with Image.open(io.BytesIO(data)) as img_check:
                    img_check.verify()

        # 3. Documento abre sem erro no python-docx e mantém os textos
        doc_after = docx.Document(out_docx)
        texts = [p.text for p in doc_after.paragraphs if p.text.strip()]
        self.assertIn("Texto inicial do parágrafo 1.", texts)
        self.assertIn("Texto intermediário entre imagens.", texts)
        self.assertIn("Texto final que não pode sumir.", texts)


if __name__ == "__main__":
    unittest.main()
