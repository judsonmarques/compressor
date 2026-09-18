"""
Módulo para alinhamento, centralização e ajuste dimensional de imagens no documento Word (.docx).
Garante que as imagens fiquem centralizadas e não estourem os limites da página nem quebrem o layout.
"""

from typing import Dict, Any, List, Optional
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, Cm
from core.config import FormattingConfig


# Prefixos comuns de legendas de imagens em português e inglês
CAPTION_PREFIXES = (
    "figura", "fig.", "imagem", "foto", "quadro", "gráfico", "grafico", "ilustração",
    "source:", "fonte:", "figure", "photo", "image", "chart"
)


def _get_max_usable_dimensions_emu(doc: Document) -> tuple[int, int]:
    """Calcula a largura e altura máximas imprimíveis da primeira seção em EMUs."""
    if not doc.sections:
        # Padrão A4 com margem 2.5cm
        width_cm = 21.0 - 5.0
        height_cm = 29.7 - 5.0
        return int(width_cm * 360000), int(height_cm * 360000)

    section = doc.sections[0]
    usable_w = (section.page_width or Cm(21.0)) - (section.left_margin or Cm(2.5)) - (section.right_margin or Cm(2.5))
    usable_h = (section.page_height or Cm(29.7)) - (section.top_margin or Cm(2.5)) - (section.bottom_margin or Cm(2.5))

    return int(usable_w), int(usable_h)


def _is_image_paragraph(p) -> bool:
    """Verifica se o parágrafo contém imagens (DrawingML ou VML pict)."""
    elements = p._p.xpath('.//w:drawing | .//w:pict')
    return len(elements) > 0


def _is_caption_paragraph(p) -> bool:
    """Verifica se o parágrafo é uma provável legenda de imagem."""
    txt = p.text.strip().lower()
    return any(txt.startswith(prefix) for prefix in CAPTION_PREFIXES)


def adjust_and_align_images(doc: Document, config: FormattingConfig) -> Dict[str, Any]:
    """
    Centraliza parágrafos com imagens e redimensiona proporcionalmente imagens que
    ultrapassem a área útil das margens da página.
    """
    images_found = 0
    images_centered = 0
    images_resized = 0

    max_w_emu, max_h_emu = _get_max_usable_dimensions_emu(doc)

    paragraphs = list(doc.paragraphs)
    num_paragraphs = len(paragraphs)

    for i, p in enumerate(paragraphs):
        if not _is_image_paragraph(p):
            continue

        images_found += 1

        # 1. Centralização da imagem
        if config.center_images:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # Remover qualquer recuo de parágrafo para a imagem não ficar deslocada
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.left_indent = Inches(0)
            p.paragraph_format.right_indent = Inches(0)

            # Espaçamento estético antes e depois
            p.paragraph_format.space_before = Pt(config.image_space_before_pt)
            p.paragraph_format.space_after = Pt(config.image_space_after_pt)

            images_centered += 1

        # 2. Vínculo com a legenda seguinte (evita imagem numa página e legenda na outra)
        if i + 1 < num_paragraphs:
            next_p = paragraphs[i + 1]
            if _is_caption_paragraph(next_p):
                p.paragraph_format.keep_with_next = True
                if config.center_images:
                    next_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    next_p.paragraph_format.first_line_indent = Inches(0)

        # 3. Redimensionamento de imagens que extrapolam as margens
        if config.fit_images_to_page:
            # Buscar elementos de extensão DrawingML (wp:extent)
            extents = p._p.xpath('.//wp:extent')
            for extent in extents:
                try:
                    cx = int(extent.get('cx', '0'))
                    cy = int(extent.get('cy', '0'))

                    if cx <= 0 or cy <= 0:
                        continue

                    needs_resize = False
                    scale = 1.0

                    # 98% da largura máxima para margem de segurança contra overflow
                    safe_max_w = int(max_w_emu * 0.98)
                    safe_max_h = int(max_h_emu * 0.90)

                    if cx > safe_max_w:
                        scale = min(scale, safe_max_w / cx)
                        needs_resize = True

                    if cy > safe_max_h:
                        scale = min(scale, safe_max_h / cy)
                        needs_resize = True

                    if needs_resize and scale < 1.0:
                        new_cx = int(cx * scale)
                        new_cy = int(cy * scale)

                        extent.set('cx', str(new_cx))
                        extent.set('cy', str(new_cy))

                        # Atualizar também elementos a:ext dentro de xfrm
                        drawing_parent = extent.getparent().getparent()
                        if drawing_parent is not None:
                            a_exts = drawing_parent.xpath('.//a:xfrm/a:ext')
                            for a_ext in a_exts:
                                a_ext.set('cx', str(new_cx))
                                a_ext.set('cy', str(new_cy))

                        images_resized += 1
                except (ValueError, TypeError):
                    continue

    # Processar também imagens embutidas em células de tabelas
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if _is_image_paragraph(p):
                        images_found += 1
                        if config.center_images:
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            p.paragraph_format.first_line_indent = Inches(0)
                            images_centered += 1

    return {
        "images_found": images_found,
        "images_centered": images_centered,
        "images_resized": images_resized,
    }
