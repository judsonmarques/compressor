"""
Módulo para eliminação de espaços em branco, parágrafos vazios e quebras redundantes.
"""

import re
from typing import Dict, Any, List
from docx import Document
from core.config import FormattingConfig


def _paragraph_has_objects(p) -> bool:
    """Verifica se o parágrafo possui imagens, desenhos ou objetos embutidos."""
    # Procura por elementos de desenho (DrawingML), VML pict, objetos OLE, formas e elementos gráficos
    elements = p._p.xpath(
        './/w:drawing | .//w:pict | .//w:object | .//*[local-name()="drawing"] | '
        './/*[local-name()="pict"] | .//*[local-name()="imagedata"] | '
        './/*[local-name()="blip"] | .//*[local-name()="shape"] | '
        './/*[local-name()="AlternateContent"] | .//w:fldSimple'
    )
    return len(elements) > 0


def _paragraph_has_page_break(p) -> bool:
    """Verifica se o parágrafo possui quebra de página explícita."""
    breaks = p._p.xpath('.//w:br[@w:type="page"]')
    return len(breaks) > 0


def clean_empty_paragraphs(doc: Document, config: FormattingConfig) -> int:
    """
    Remove parágrafos em branco excessivos e vazios do corpo do documento.
    
    Respeita max_consecutive_empty_lines (por padrão 0, eliminando todas as linhas vazias
    que costumam desalinhar o layout entre páginas).
    """
    if not config.remove_empty_paragraphs:
        return 0

    removed_count = 0
    consecutive_empty = 0
    paragraphs_to_remove = []

    # Iterar pelos parágrafos principais
    for p in doc.paragraphs:
        text = p.text.strip()
        has_objects = _paragraph_has_objects(p)
        has_page_break = _paragraph_has_page_break(p)

        is_empty = (text == "") and not has_objects and not has_page_break

        if is_empty:
            consecutive_empty += 1
            if consecutive_empty > config.max_consecutive_empty_lines:
                paragraphs_to_remove.append(p)
        else:
            consecutive_empty = 0

    # Remover os parágrafos identificados
    for p in paragraphs_to_remove:
        parent = p._p.getparent()
        if parent is not None:
            parent.remove(p._p)
            removed_count += 1

    # Remover parágrafos vazios no final do documento se configurado
    if config.remove_trailing_empty_paragraphs:
        while doc.paragraphs:
            last_p = doc.paragraphs[-1]
            if (last_p.text.strip() == "" 
                and not _paragraph_has_objects(last_p) 
                and not _paragraph_has_page_break(last_p)):
                parent = last_p._p.getparent()
                if parent is not None:
                    parent.remove(last_p._p)
                    removed_count += 1
            else:
                break

    return removed_count


def clean_redundant_spaces(doc: Document, config: FormattingConfig) -> int:
    """
    Normaliza múltiplos espaços em branco consecutivos nos textos (ex: 'palavra    outra' -> 'palavra outra').
    Preserva a formatação de fontes e estilos dos runs.
    """
    if not config.remove_redundant_spaces:
        return 0

    cleaned_runs = 0
    space_pattern = re.compile(r'[ \t]{2,}')

    # Processar parágrafos
    for p in doc.paragraphs:
        for run in p.runs:
            if run.text and space_pattern.search(run.text):
                run.text = space_pattern.sub(' ', run.text)
                cleaned_runs += 1

    # Processar tabelas
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        if run.text and space_pattern.search(run.text):
                            run.text = space_pattern.sub(' ', run.text)
                            cleaned_runs += 1

    return cleaned_runs


def clean_document_whitespace(doc: Document, config: FormattingConfig) -> Dict[str, Any]:
    """Executa a rotina completa de higienização e limpeza do documento."""
    empty_removed = clean_empty_paragraphs(doc, config)
    spaces_cleaned = clean_redundant_spaces(doc, config)

    return {
        "empty_paragraphs_removed": empty_removed,
        "runs_spaces_cleaned": spaces_cleaned,
    }
