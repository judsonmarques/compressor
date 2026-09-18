"""
Gerador de documento de teste propositalmente despadronizado e bagunçado.
Contém linhas vazias repetidas, espaços duplos, imagem gigante desalinhada,
tabela sem centralização e títulos sem controle de paginação.
"""

import os
import io
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def create_dummy_image(filename: str, width: int = 1400, height: int = 800) -> str:
    """Cria uma imagem de teste de alta resolução."""
    img = Image.new("RGB", (width, height), color=(52, 152, 219))
    draw = ImageDraw.Draw(img)
    
    # Desenhar linhas decorativas
    draw.rectangle([20, 20, width - 20, height - 20], outline=(255, 255, 255), width=8)
    
    os.makedirs(os.path.dirname(filename), exist_ok=True)
    img.save(filename, format="PNG")
    return filename


def generate_messy_docx(output_path: str):
    """Gera um arquivo .docx com múltiplos defeitos de formatação comuns."""
    doc = Document()

    # Título do Documento
    h1 = doc.add_heading("Relatório Técnico de Auditoria Operacional", level=1)
    
    # Linhas vazias redundantes no início
    doc.add_paragraph("")
    doc.add_paragraph("")
    doc.add_paragraph("   ")

    # Parágrafo 1 com múltiplos espaços duplicados e triplados
    p1 = doc.add_paragraph(
        "Este documento foi elaborado    para fins de avaliação    de desempenho dos processos internos. "
        "Nota-se que existem múltiplos      espaços em branco entre palavras que foram inseridos incorretamente "
        "durante a digitação manual dos relatórios setoriais."
    )
    p1.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Mais parágrafos vazios excessivos
    doc.add_paragraph("")
    doc.add_paragraph("")
    doc.add_paragraph("")

    # Título 2
    h2 = doc.add_heading("1. Dados e Indicadores Gerais", level=2)

    p2 = doc.add_paragraph(
        "Abaixo é apresentada a imagem analítica do setor produtivo. Observe que esta imagem originalmente "
        "foi inserida com alinhamento torto à esquerda e dimensões que ultrapassam os limites convencionais "
        "de impressão da folha A4, além de conter recuos indesejados."
    )
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Imagem de Teste (grande, ocupando largura excessiva)
    img_path = os.path.join(os.path.dirname(output_path), "sample_chart.png")
    create_dummy_image(img_path, width=1600, height=900)

    # Parágrafo com imagem desalinhada à esquerda com recuo estranho
    img_p = doc.add_paragraph()
    img_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    img_p.paragraph_format.first_line_indent = Inches(1.5)
    img_run = img_p.add_run()
    # Inserir propositalmente com 8.5 polegadas (ultrapassa a margem da folha A4!)
    img_run.add_picture(img_path, width=Inches(8.5))

    # Legenda da imagem separada
    cap_p = doc.add_paragraph("Figura 1 - Visão geral do fluxo operacional e indicadores de produtividade.")
    cap_p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Linhas vazias consecutivas
    doc.add_paragraph("")
    doc.add_paragraph("")
    doc.add_paragraph("")

    # Tabela não centralizada e sem proteção de quebra de página
    h3 = doc.add_heading("2. Detalhamento Quantitativo", level=2)
    tbl = doc.add_table(rows=4, cols=3)
    
    headers = ["Indicador", "Meta Estabelecida", "Resultado Obtido"]
    for j, h in enumerate(headers):
        tbl.cell(0, j).text = h

    data = [
        ["Taxa de Conformidade", "95.0%", "98.4%"],
        ["Tempo Médio de Atendimento", "< 15 min", "11.2 min"],
        ["Índice de Retrabalho", "< 2.0%", "0.8%"],
    ]
    for i, row in enumerate(data, start=1):
        for j, val in enumerate(row):
            tbl.cell(i, j).text = val

    # Linhas vazias no fim do documento
    doc.add_paragraph("")
    doc.add_paragraph("")
    doc.add_paragraph("")

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Documento de teste gerado em: {output_path}")


if __name__ == "__main__":
    test_docx = os.path.abspath(os.path.join(os.path.dirname(__file__), "documento_exemplo_desalinhado.docx"))
    generate_messy_docx(test_docx)
