"""
Testes de verificação da nova funcionalidade de redução de tamanho de arquivos.
Valida compressão de PDF, DOCX e Imagens com garantia de integridade.
"""

import os
import io
import unittest
from PIL import Image, ImageDraw
import pymupdf
from docx import Document

from core.compressor import FileCompressor, CompressionConfig, CompressionResult


class TestFileCompressor(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.test_dir = os.path.join(os.path.dirname(__file__), "temp_compression_tests")
        os.makedirs(cls.test_dir, exist_ok=True)

        # 1. Criar imagem grande de teste
        cls.img_path = os.path.join(cls.test_dir, "imagem_teste_grande.jpg")
        img = Image.new("RGB", (2000, 1500), color=(80, 120, 200))
        draw = ImageDraw.Draw(img)
        for i in range(0, 2000, 100):
            draw.line([(i, 0), (i, 1500)], fill=(255, 255, 255), width=5)
        img.save(cls.img_path, format="JPEG", quality=95)

        # 2. Criar PDF de teste contendo imagem grande
        cls.pdf_path = os.path.join(cls.test_dir, "documento_teste.pdf")
        doc_pdf = pymupdf.open()
        page = doc_pdf.new_page(width=595, height=842)
        page.insert_text((50, 50), "Relatório com Imagem para Teste de Redução de Tamanho", fontsize=14)
        with open(cls.img_path, "rb") as f:
            img_bytes = f.read()
        page.insert_image(pymupdf.Rect(50, 100, 545, 450), stream=img_bytes)
        doc_pdf.save(cls.pdf_path)
        doc_pdf.close()

        # 3. Criar DOCX de teste contendo imagem grande
        cls.docx_path = os.path.join(cls.test_dir, "documento_teste.docx")
        doc_docx = Document()
        doc_docx.add_heading("Documento DOCX com Imagem Pesada", level=1)
        doc_docx.add_paragraph("Este parágrafo de texto deve permanecer 100% íntegro.")
        doc_docx.add_picture(cls.img_path)
        doc_docx.save(cls.docx_path)

    @classmethod
    def tearDownClass(cls):
        # Limpar arquivos temporários
        if os.path.exists(cls.test_dir):
            for f in os.listdir(cls.test_dir):
                try:
                    os.remove(os.path.join(cls.test_dir, f))
                except Exception:
                    pass
            try:
                os.rmdir(cls.test_dir)
            except Exception:
                pass

    def test_compress_pdf(self):
        compressor = FileCompressor()
        out_pdf = os.path.join(self.test_dir, "documento_teste_reduzido.pdf")
        
        result = compressor.compress_file(self.pdf_path, output_path=out_pdf)
        self.assertTrue(result.success, f"Falha na compressão do PDF: {result.error_message}")
        self.assertTrue(os.path.exists(out_pdf))
        self.assertLess(result.compressed_size, result.original_size)
        self.assertGreater(result.reduction_percentage, 10.0)
        
        # Validar integridade do PDF gerado
        doc = pymupdf.open(out_pdf)
        self.assertEqual(len(doc), 1)
        text = doc[0].get_text()
        self.assertIn("Relatório", text)
        doc.close()

    def test_compress_docx(self):
        compressor = FileCompressor()
        out_docx = os.path.join(self.test_dir, "documento_teste_reduzido.docx")

        result = compressor.compress_file(self.docx_path, output_path=out_docx)
        self.assertTrue(result.success, f"Falha na compressão do DOCX: {result.error_message}")
        self.assertTrue(os.path.exists(out_docx))
        self.assertLess(result.compressed_size, result.original_size)

        # Validar integridade do DOCX gerado abrindo com python-docx
        doc = Document(out_docx)
        self.assertEqual(doc.paragraphs[0].text, "Documento DOCX com Imagem Pesada")
        self.assertEqual(doc.paragraphs[1].text, "Este parágrafo de texto deve permanecer 100% íntegro.")

    def test_compress_image(self):
        compressor = FileCompressor()
        out_img = os.path.join(self.test_dir, "imagem_teste_reduzida.jpg")

        result = compressor.compress_file(self.img_path, output_path=out_img)
        self.assertTrue(result.success)
        self.assertTrue(os.path.exists(out_img))
        self.assertLess(result.compressed_size, result.original_size)

    def test_presets(self):
        presets = CompressionConfig.get_presets()
        self.assertIn("Alta Fidelidade (Mínima Perda)", presets)
        self.assertIn("Equilibrado (Recomendado)", presets)
        self.assertIn("Máxima Redução", presets)


if __name__ == "__main__":
    unittest.main()
